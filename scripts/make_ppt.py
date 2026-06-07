#!/usr/bin/env python3
# 선행기구개발그룹 시료 수량 취합 결과 → PPTX 생성
import json, urllib.request, datetime
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

DB = "https://test-set-check-default-rtdb.asia-southeast1.firebasedatabase.app/submissions.json"

PARTS = [
    ("ne", "NE파트"), ("package", "패키지파트"), ("cmf1", "선행CMF1파트"),
    ("cmf2", "선행CMF2파트"), ("cmf3", "선행CMF3파트"), ("hinge", "HINGE개발LAB"),
    ("bonding", "접합기술파트"),
]
GROUPS = [
    ("strategic_hhp", "전략HHP"), ("innovative_hhp", "혁신HHP"), ("tablet", "태블릿"),
    ("npc", "NPC"), ("watch", "WATCH"), ("tws", "TWS"), ("xr", "XR"), ("glass", "GLASS"),
]
BASELINE = 20
KFONT = "맑은 고딕"  # Malgun Gothic

# ---- 데이터 로드 ----
with urllib.request.urlopen(DB, timeout=20) as r:
    data = json.load(r) or {}

def cell(pk, gk):
    return (data.get(pk) or {}).get(gk)

def qty(pk, gk):
    c = cell(pk, gk)
    if not c or c.get("quantity") is None:
        return None
    return int(c["quantity"])

def part_reason(pk):
    """파트의 제품군별 사유를 동일 사유끼리 묶어서 정리."""
    by_reason = {}
    author = ""
    for gk, glabel in GROUPS:
        c = cell(pk, gk)
        if not c:
            continue
        author = c.get("author") or author
        rsn = (c.get("reason") or "").strip()
        if not rsn:
            continue
        by_reason.setdefault(rsn, []).append(glabel)
    lines = []
    for rsn, glabels in by_reason.items():
        tag = "·".join(glabels)
        lines.append(f"[{tag}] {rsn}")
    return author, lines

# ---- 프레젠테이션 ----
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
HEADER = RGBColor(0x2F, 0x6D, 0xF6)
SUMBG = RGBColor(0x1F, 0x3A, 0x5F)
ZEBRA = RGBColor(0xF2, 0xF5, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1C, 0x23, 0x33)
GREEN = RGBColor(0xE3, 0xF6, 0xEC)
BLUE = RGBColor(0xE4, 0xEE, 0xFE)
ORANGE = RGBColor(0xFD, 0xEA, 0xE0)

def set_cell(tc, text, *, size=11, bold=False, color=DARK, fill=None,
             align=PP_ALIGN.CENTER, font=KFONT):
    tc.margin_left = Inches(0.05); tc.margin_right = Inches(0.05)
    tc.margin_top = Inches(0.02); tc.margin_bottom = Inches(0.02)
    tc.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill is not None:
        tc.fill.solid(); tc.fill.fore_color.rgb = fill
    tf = tc.text_frame; tf.word_wrap = True
    lines = str(text).split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run(); run.text = ln
        run.font.size = Pt(size); run.font.bold = bold
        run.font.name = run.font.name = KFONT
        run.font.color.rgb = color

# =========================================================================
# Slide 1 — 표지
# =========================================================================
s = prs.slides.add_slide(blank)
bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
bg.shadow.inherit = False
s.shapes._spTree.remove(bg._element); s.shapes._spTree.insert(2, bg._element)

def textbox(slide, l, t, w, h, text, size, bold, color, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run(); run.text = ln
        run.font.size = Pt(size); run.font.bold = bold
        run.font.name = KFONT; run.font.color.rgb = color
    return tb

textbox(s, 0.9, 2.3, 11.5, 1.2, "선행기구개발그룹", 24, True, RGBColor(0x9D, 0xC2, 0xFF))
textbox(s, 0.9, 3.0, 11.5, 1.6,
        "모델별 양산단계 시료 보유 수량 취합 결과", 40, True, WHITE)
textbox(s, 0.9, 4.7, 11.5, 0.8,
        f"디폴트 기준: 모델별 {BASELINE}대 (팀 전체) · 파트별 실제 수량 취합 후 조율", 16, False,
        RGBColor(0xCF, 0xDC, 0xF5))
today = datetime.date.today().strftime("%Y-%m-%d")
textbox(s, 0.9, 6.4, 11.5, 0.6, f"작성일: {today}", 14, False, RGBColor(0xAE, 0xBE, 0xD9))

# =========================================================================
# Slide 2 — 취합 표
# =========================================================================
s = prs.slides.add_slide(blank)
textbox(s, 0.4, 0.25, 12.5, 0.7,
        "파트별 · 제품군별 제안 시료 수량", 24, True, NAVY)
textbox(s, 0.42, 0.92, 12.5, 0.4,
        f"(단위: 대 / 디폴트 기준 {BASELINE}대 · 빈칸은 미입력)", 12, False, RGBColor(0x67, 0x70, 0x8A))

ncols = 1 + len(GROUPS) + 1          # 파트 + 제품군8 + 사유
nrows = 1 + len(PARTS) + 1           # 헤더 + 파트7 + 합계

left, top = Inches(0.35), Inches(1.4)
width = Inches(12.65); height = Inches(5.6)
gtab = s.shapes.add_table(nrows, ncols, left, top, width, height).table

# 열 너비
gtab.columns[0].width = Inches(1.35)            # 파트
for i in range(len(GROUPS)):
    gtab.columns[1 + i].width = Inches(0.82)    # 제품군
gtab.columns[ncols - 1].width = Inches(4.74)    # 사유

# 헤더 행
set_cell(gtab.cell(0, 0), "파트 \\ 제품군", size=11, bold=True, color=WHITE, fill=HEADER)
for i, (gk, gl) in enumerate(GROUPS):
    set_cell(gtab.cell(0, 1 + i), gl, size=10.5, bold=True, color=WHITE, fill=HEADER)
set_cell(gtab.cell(0, ncols - 1), "사유", size=11, bold=True, color=WHITE, fill=HEADER)

col_tot = [0] * len(GROUPS)

# 파트 행
for r, (pk, pl) in enumerate(PARTS, start=1):
    author, rlines = part_reason(pk)
    pname = pl + (f"\n({author})" if author else "")
    rowfill = ZEBRA if (r % 2 == 0) else WHITE
    set_cell(gtab.cell(r, 0), pname, size=10.5, bold=True, color=NAVY, fill=rowfill)
    for i, (gk, gl) in enumerate(GROUPS):
        q = qty(pk, gk)
        if q is None:
            set_cell(gtab.cell(r, 1 + i), "", size=10, fill=rowfill)
        else:
            col_tot[i] += q
            if q == 0:
                fill = rowfill
            elif q == BASELINE:
                fill = GREEN
            elif q < BASELINE:
                fill = BLUE
            else:
                fill = ORANGE
            set_cell(gtab.cell(r, 1 + i), str(q), size=11, bold=True, fill=fill)
    reason_text = "\n".join(rlines) if rlines else "-"
    set_cell(gtab.cell(r, ncols - 1), reason_text, size=8, color=DARK,
             fill=rowfill, align=PP_ALIGN.LEFT)

# 합계 행
rr = nrows - 1
set_cell(gtab.cell(rr, 0), "합계", size=11, bold=True, color=WHITE, fill=SUMBG)
for i in range(len(GROUPS)):
    set_cell(gtab.cell(rr, 1 + i), str(col_tot[i]), size=11, bold=True, color=WHITE, fill=SUMBG)
set_cell(gtab.cell(rr, ncols - 1), f"총 {sum(col_tot)}대", size=10, bold=True,
         color=WHITE, fill=SUMBG, align=PP_ALIGN.LEFT)

# 행 높이
gtab.rows[0].height = Inches(0.4)
for r in range(1, nrows - 1):
    gtab.rows[r].height = Inches(0.72)
gtab.rows[nrows - 1].height = Inches(0.4)

out = "선행기구개발그룹_시료수량_취합결과.pptx"
prs.save(out)
print("SAVED:", out)
print("열 수:", ncols, "행 수:", nrows)
print("제품군별 합계:", dict(zip([g[1] for g in GROUPS], col_tot)))
print("총합:", sum(col_tot))
